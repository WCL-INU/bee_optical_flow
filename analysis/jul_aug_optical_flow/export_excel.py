"""Export original July/August flow values plus approved linear-model predictions."""
from datetime import datetime
from pathlib import Path
import hashlib
import json
import math
import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from count_calibration import apply_models, DERIVED_COLUMNS

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
from src.edit_roi import parse_boundary, read_records

RESULTS = ROOT / 'bee_count_output/yaml_jul_aug_all'
SOURCE = RESULTS / 'batch_summary.csv'
OUTPUT = ROOT / 'optical_flow_jul_aug_analysis_data.xlsx'
PATTERN = re.compile(r'ANU-25-summer-(\d+)_(20260[78]\d{2})_(\d{6})\.mp4$')

# All exported measurement columns come directly from batch_summary.csv.
DEFINITIONS = [
    ('video', '영상 정보', '파일명', '기기 번호와 촬영 일시가 포함된 원본 영상 파일명'),
    ('video_path', '영상 정보', '경로', '원본 CSV에 기록된 영상 경로'),
    ('duration_sec', '영상 정보', '초', 'optical flow 집계에 사용된 처리 구간의 시간'),
    ('processed_frame_pairs', '영상 정보', '쌍', 'optical flow를 계산한 연속 프레임 쌍 수'),
]
for prefix, category, description in [
    ('roi', '합산 영역', '전체 영상 기준 optical flow 계산 ROI'),
    ('ent', '합산 영역', '전체 영상 기준 입구 사각형'),
    ('entrance_roi', '합산 영역', 'ROI 내부 상대 좌표의 입구 사각형'),
]:
    for coordinate in ['x1', 'y1', 'x2', 'y2']:
        DEFINITIONS.append((f'{prefix}_{coordinate}', category, '픽셀', f'{description}: {coordinate}'))
DEFINITIONS.append(('boundary_band_px', '합산 영역', '픽셀', '입구 경계 띠를 구성하는 폭 설정; 실제 합산 픽셀 수와 구별'))
for mode, label in [('raw', '후보 픽셀, persistence·component 면적 필터 적용 전'),
                    ('filtered', '필터를 통과한 픽셀')]:
    for direction, description in [('in', '입구 안쪽 방향'), ('out', '입구 바깥쪽 방향'),
                                    ('traffic', 'IN + OUT')]:
        DEFINITIONS.append((f'total_{mode}_{direction}_flux', 'Optical flow', 'flux',
                            f'{description} 법선 flow의 영상 전체 누적 합; {label}'))
    DEFINITIONS.append((f'mean_{mode}_traffic_flux_per_sec', 'Optical flow', 'flux/초',
                        f'total_{mode}_traffic_flux / duration_sec'))
DEFINITIONS.extend([
    ('raw_to_filtered_reduction_ratio', 'Optical flow', '배율',
     'raw traffic / max(filtered traffic, epsilon); 제거율(%) 아님. 원본 코드의 epsilon은 1e-6'),
    ('frame_csv', '결과 파일', '경로', '프레임 쌍별 optical flow 결과 CSV 경로'),
    ('window_csv', '결과 파일', '경로', '3초 구간별 optical flow 결과 CSV 경로'),
])
SOURCE_COLUMNS = [item[0] for item in DEFINITIONS]
for direction, label in [('in', 'IN'), ('out', 'OUT'), ('traffic', 'IN + OUT')]:
    expression = ('방향별 기울기 × total_filtered_' + direction + '_flux + 절편'
                  if direction != 'traffic' else 'linear_predicted_in_count + linear_predicted_out_count')
    DEFINITIONS.append((f'linear_predicted_{direction}_count', '선형 회귀 추정', '추정 마리/영상',
                        f'{label} 영상별 추정 마리수: {expression}; 원본 코드의 flux/100 count_est와 구별'))
for direction, label in [('in', 'IN'), ('out', 'OUT'), ('traffic', 'IN + OUT')]:
    DEFINITIONS.append((f'linear_predicted_{direction}_count_per_min', '선형 회귀 추정', '추정 마리/분',
                        f'{label} 분당 추정 마리수 = linear_predicted_{direction}_count / duration_sec × 60'))
COLUMNS = [item[0] for item in DEFINITIONS]
FLOW_COLUMNS = [item[0] for item in DEFINITIONS if item[1] == 'Optical flow']


def main():
    batch = pd.read_csv(SOURCE, float_precision='round_trip')
    skipped = pd.read_csv(RESULTS / 'skipped_videos.csv')
    calibration = json.loads((HERE / 'linear_count_calibration.json').read_text())
    assert calibration['source_csv_sha256'] == hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    for key in ['model_source', 'training_source']:
        assert calibration['metadata'][f'{key}_sha256'] == hashlib.sha256(
            (ROOT / calibration['metadata'][key]).read_bytes()).hexdigest()
    selected = apply_models(batch, calibration['metadata'])[COLUMNS]
    summary_names = set(batch.video)
    source_names = {entry.name for entry in os.scandir(ROOT / 'videos')
                    if PATTERN.fullmatch(entry.name) and entry.is_file()}
    entries = [entry.name for entry in os.scandir(RESULTS)]
    frames = {name.removesuffix('_frame_flux.csv') + '.mp4'
              for name in entries if name.endswith('_frame_flux.csv')}
    windows = {name.removesuffix('_window_3sec.csv') + '.mp4'
               for name in entries if name.endswith('_window_3sec.csv')}
    skip_names = set(skipped.video)
    periods = [(record, parse_boundary(str(record['start'])), parse_boundary(str(record['end']), True))
               for record in read_records(ROOT / 'roi_regions.yaml')['regions']]
    eligible = set()
    records = []
    for name in sorted(source_names | summary_names | frames | windows | skip_names):
        match = PATTERN.fullmatch(name)
        assert match, name
        device = f'ANU-25-summer-{match[1]}'
        stamp = datetime.strptime(match[2] + match[3], '%Y%m%d%H%M%S')
        matches = [record for record, start, end in periods
                   if record['device'] == device and start <= stamp <= end]
        if name in source_names and len(matches) == 1:
            eligible.add(name)
        records.append(dict(video=name, device=device, recorded_at=stamp, month=stamp.strftime('%Y-%m'),
                            source_present=name in source_names, batch_present=name in summary_names,
                            frame_present=name in frames, window_present=name in windows,
                            yaml_period_matches=len(matches), skipped_listed=name in skip_names))
    inventory = pd.DataFrame(records)
    monthly = []
    for month, group in inventory.groupby('month'):
        monthly.append(dict(month=month, source_videos=int(group.source_present.sum()),
                            yaml_eligible=int((group.source_present & group.yaml_period_matches.eq(1)).sum()),
                            batch_videos=int(group.batch_present.sum()),
                            frame_results=int(group.frame_present.sum()), window_results=int(group.window_present.sum()),
                            skipped_videos=int(group.skipped_listed.sum()),
                            unaccounted_videos=int((group.source_present & ~group.batch_present & ~group.skipped_listed).sum())))
    missing_paths = []
    for column, suffix in [('frame_csv', '_frame_flux.csv'), ('window_csv', '_window_3sec.csv')]:
        for video, value in zip(batch.video, batch[column]):
            path = Path(value)
            if not path.is_absolute():
                path = ROOT / path
            if not path.is_file() or path.name != Path(video).stem + suffix:
                missing_paths.append(dict(video=video, column=column, path=value))
    audit = dict(
        source_csv=str(SOURCE.relative_to(ROOT)), source_csv_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        workbook=OUTPUT.name, selected_columns=COLUMNS, selected_column_count=len(COLUMNS),
        source_videos=len(source_names), batch_rows=len(batch), batch_unique_videos=len(summary_names),
        duplicate_rows=int(batch.video.duplicated().sum()), yaml_eligible_videos=len(eligible),
        frame_results=len(frames), window_results=len(windows), skipped_videos=len(skip_names),
        source_without_batch=sorted(source_names - summary_names),
        unaccounted_source_videos=sorted(source_names - summary_names - skip_names),
        eligible_without_batch=sorted(eligible - summary_names),
        batch_without_source=sorted(summary_names - source_names),
        batch_without_frames=sorted(summary_names - frames), batch_without_windows=sorted(summary_names - windows),
        frames_without_batch=sorted(frames - summary_names), windows_without_batch=sorted(windows - summary_names),
        missing_or_mismatched_result_paths=missing_paths,
        flow_missing_cells=int(selected[FLOW_COLUMNS].isna().sum().sum()),
        flow_nonfinite_cells=int((~np.isfinite(selected[FLOW_COLUMNS].to_numpy())).sum()),
        geometry_blank_rows=int(selected.roi_x1.isna().sum()), monthly=monthly,
        original_column_count=len(SOURCE_COLUMNS), derived_columns=DERIVED_COLUMNS,
        derived_missing_cells=int(selected[DERIVED_COLUMNS].isna().sum().sum()),
        calibration_metadata=calibration['metadata'],
        excluded_columns=[column for column in batch.columns if column not in COLUMNS],
    )
    assert audit['duplicate_rows'] == 0
    assert source_names == summary_names | skip_names and not summary_names & skip_names
    assert summary_names == eligible == frames == windows
    assert not missing_paths and audit['flow_missing_cells'] == audit['flow_nonfinite_cells'] == 0
    assert audit['derived_missing_cells'] == 0
    assert batch.video.map(lambda name: bool(PATTERN.fullmatch(name))).all()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Optical flow'
    ws.append(COLUMNS)
    for row in selected.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value for value in row])
    for index, definition in enumerate(DEFINITIONS, 1):
        ws.cell(1, index).comment = Comment(f'{definition[1]} | {definition[2]}\n{definition[3]}', 'Column dictionary')
        ws.column_dimensions[get_column_letter(index)].width = 46 if definition[0] == 'video' else (
            58 if definition[0].endswith('csv') or definition[0] == 'video_path' else
            34 if definition[0] in DERIVED_COLUMNS else 25 if 'flux' in definition[0] else 20)
        if pd.api.types.is_numeric_dtype(selected[definition[0]]):
            fmt = '#,##0' if 'flux' not in definition[0] and definition[0] not in (
                'duration_sec', 'raw_to_filtered_reduction_ratio', *DERIVED_COLUMNS) else '#,##0.000000'
            for cells in ws.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in cells:
                    cell.number_format = fmt
    ws.freeze_panes = 'C2'
    ws.sheet_view.zoomScale = 75
    data_table = Table(displayName='OpticalFlowResults', ref=ws.dimensions)
    data_table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(data_table)

    dictionary = wb.create_sheet('컬럼 설명')
    dictionary.append(['컬럼', '분류', '단위', '설명', '결측 행 수'])
    for definition in DEFINITIONS:
        dictionary.append([*definition, int(selected[definition[0]].isna().sum())])
    for index, width in enumerate([44, 20, 16, 110, 20], 1):
        dictionary.column_dimensions[get_column_letter(index)].width = width
    dictionary.auto_filter.ref = dictionary.dimensions

    check = wb.create_sheet('포함 여부 확인')
    check.append(['확인 항목', '결과', '설명'])
    checks = [
        ('원본 7–8월 영상', len(source_names), 'videos 디렉토리 실제 파일명 목록'),
        ('YAML 기간에 포함된 원본 영상', len(eligible), '현행 ROI YAML의 기기·촬영 일시와 대조'),
        ('batch_summary 영상', len(summary_names), '영상명 중복 없음; 엑셀에 같은 순서로 모든 행 수록'),
        ('프레임 결과 CSV', len(frames), 'batch_summary 영상명 집합과 동일'),
        ('3초 구간 결과 CSV', len(windows), 'batch_summary 영상명 집합과 동일'),
        ('추출 대상 누락', len(eligible - summary_names), 'YAML에 포함되고 원본이 있는 영상 중 batch_summary 누락'),
        ('제외 기록 영상', len(skip_names), '8월 10일 영상 80개; YAML 기간 공백과 skipped_videos.csv 일치'),
        ('설명되지 않은 원본 영상', len(source_names - summary_names - skip_names), '원본에서 결과와 제외 기록 양쪽에 없는 영상'),
        ('영상명 중복 행', audit['duplicate_rows'], 'batch_summary.csv 원본 기준'),
        ('핵심 flow 컬럼 결측 셀', audit['flow_missing_cells'], '0 신호도 정상 수치로 유지'),
        ('엑셀 컬럼 수', len(COLUMNS), '원본 28개 컬럼 + 선형 모델 추정치 6개 컬럼'),
        ('ROI/ENT 설정 공란 행', audit['geometry_blank_rows'], '원본 CSV 공란을 그대로 보존; 외부 YAML로 채우지 않음'),
        ('수치 저장', '원본 + 선형 회귀 추정', '원본 값 보존, 면적 환산 없음; Excel 숫자는 약 15자리 정밀도'),
        ('추정치 연산', '기울기 × 방향별 flux + 절편', '영상마다 방향별 절편을 한 번 적용; 반올림·클리핑 없음'),
        ('원본 CSV', str(SOURCE.relative_to(ROOT)), '동일 원본의 SHA-256을 아래 기록'),
        ('원본 CSV SHA-256', audit['source_csv_sha256'], '대조한 CSV 파일의 해시'),
    ]
    for item in checks:
        check.append(item)
    for index, width in enumerate([38, 74, 105], 1):
        check.column_dimensions[get_column_letter(index)].width = width
    period_sheet = wb.create_sheet('월별 포함 현황')
    period_sheet.append(['월', '원본 영상', 'YAML 포함', 'batch_summary', '프레임 결과', '3초 구간 결과', '제외 목록', '미확인 원본'])
    for item in monthly:
        period_sheet.append(list(item.values()))
    for index in range(1, 9):
        period_sheet.column_dimensions[get_column_letter(index)].width = 20
    skipped_sheet = wb.create_sheet('추출 제외 영상')
    skipped_sheet.append(['video', 'video_path', 'device', 'recorded_at', 'reason'])
    for row in skipped[['video', 'video_path', 'device', 'recorded_at', 'reason']].itertuples(index=False, name=None):
        skipped_sheet.append(row)
    for index, width in enumerate([48, 100, 25, 25, 85], 1):
        skipped_sheet.column_dimensions[get_column_letter(index)].width = width
    skipped_sheet.auto_filter.ref = skipped_sheet.dimensions
    model_sheet = wb.create_sheet('선형 회귀 모델')
    model_sheet.append(['방향', '입력 컬럼', '기울기', '절편', '학습 표본', 'R²', '학습 MAE', '학습 RMSE', '예측식'])
    for direction in ['in', 'out']:
        model = calibration['metadata']['models'][direction]
        model_sheet.append([direction.upper(), model['x_col'], model['slope'], model['intercept'], model['n'],
                            model['r_squared'], model['mae'], model['rmse'],
                            f"{model['slope']:.17g} × {model['x_col']} + {model['intercept']:.17g}"])
    model_sheet.append([])
    for key, label in [('model_source', '계수 출처'), ('model_source_sha256', '계수 CSV SHA-256'),
                       ('training_source', '학습 자료'), ('training_source_sha256', '학습 자료 SHA-256'),
                       ('training_start', '학습 시작'), ('training_end', '학습 끝')]:
        model_sheet.append([label, calibration['metadata'][key]])
    model_sheet.append(['절편 적용', '영상마다 방향별 1회; flux=0에서도 절편 유지'])
    model_sheet.append(['합계 추정', 'IN 예측 + OUT 예측'])
    model_sheet.append(['분당 추정', '영상별 예측 / duration_sec × 60'])
    model_sheet.append(['수치 정책', '반올림·클리핑·면적 정규화 없음'])
    for index, width in enumerate([24, 72, 26, 26, 18, 18, 20, 20, 98], 1):
        model_sheet.column_dimensions[get_column_letter(index)].width = width
    for row in range(2, 4):
        model_sheet.cell(row, 3).number_format = '0.000000000000E+00'
        for col in [4, 6, 7, 8]:
            model_sheet.cell(row, col).number_format = '0.000000000000'
    for sheet in wb:
        if sheet != ws:
            sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 36
    wb.save(OUTPUT)

    # Verify the actual written workbook, with allowance only for Excel's numeric precision.
    written = load_workbook(OUTPUT, read_only=True, data_only=True)
    sheet = written['Optical flow']
    assert sheet.max_row == len(batch) + 1 and sheet.max_column == len(COLUMNS)
    rows = sheet.iter_rows(values_only=True)
    assert list(next(rows)) == COLUMNS
    checked = 0
    for expected, actual in zip(selected.itertuples(index=False, name=None), rows, strict=True):
        for source_value, excel_value in zip(expected, actual, strict=True):
            if pd.isna(source_value):
                assert excel_value is None
            elif isinstance(source_value, (int, float)):
                assert isinstance(excel_value, (int, float))
                assert math.isclose(source_value, excel_value, rel_tol=1e-14, abs_tol=1e-12)
            else:
                assert source_value == excel_value
            checked += 1
    assert written['추출 제외 영상'].max_row == len(skipped) + 1
    assert written['선형 회귀 모델'].cell(2, 5).value == calibration['metadata']['models']['in']['n']
    assert len(written.sheetnames) == 6
    written.close()
    audit['verified_export_cells'] = checked
    audit['workbook_verified'] = True
    (HERE / 'batch_summary_coverage.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n')
    inventory.to_csv(HERE / 'tables/batch_summary_coverage.csv', index=False, encoding='utf-8-sig')
    print(json.dumps({key: audit[key] for key in ['workbook', 'source_videos', 'batch_rows', 'skipped_videos',
                                                 'selected_column_count', 'monthly', 'verified_export_cells']},
                     ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
